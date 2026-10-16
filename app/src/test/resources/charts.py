# Chart kinds test workbook: a combo (columns and a line), a radar, a bubble chart — one sheet each.
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, RadarChart, BubbleChart, Reference, Series
wb = Workbook()
ws = wb.active; ws.title = "Combo"
ws.append(["Month", "Sales", "Target"])
for row in [["Jan", 120, 100], ["Feb", 90, 110], ["Mar", 150, 120], ["Apr", 130, 130], ["May", 170, 140]]: ws.append(row)
bar = BarChart(); bar.title = "Sales and target"
bar.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
line = LineChart(); line.add_data(Reference(ws, min_col=3, min_row=1, max_row=6), titles_from_data=True)
bar += line
ws.add_chart(bar, "E2")
r = wb.create_sheet("Radar")
r.append(["Skill", "Anna", "Nikos"])
for row in [["Reading", 8, 6], ["Writing", 6, 9], ["Maths", 9, 5], ["Science", 7, 8], ["Art", 5, 7]]: r.append(row)
rc = RadarChart(); rc.type = "marker"; rc.title = "Skills"
rc.add_data(Reference(r, min_col=2, max_col=3, min_row=1, max_row=6), titles_from_data=True)
rc.set_categories(Reference(r, min_col=1, min_row=2, max_row=6))
r.add_chart(rc, "E2")
b = wb.create_sheet("Bubble")
b.append(["Price", "Sold", "Shops"])
for row in [[10, 50, 3], [20, 35, 8], [30, 42, 5], [40, 20, 12]]: b.append(row)
bc = BubbleChart(); bc.title = "Price and sales"
s = Series(values=Reference(b, min_col=2, min_row=2, max_row=5), xvalues=Reference(b, min_col=1, min_row=2, max_row=5),
           zvalues=Reference(b, min_col=3, min_row=2, max_row=5), title="Shops")
bc.series.append(s)
b.add_chart(bc, "E2")
wb.save(sys.argv[1])
