# cropped.docx / cropped.xlsx / cropped.pptx: one 900x600 picture (green sides, a black band on top, an orange sun) cropped
# 2/9 off each side and 1/10 off the top, so only the sun on white shows (as LibreOffice draws it).
# cropped.odp, cropped.odt: LibreOffice's own (soffice --headless --convert-to odp cropped.pptx; --convert-to odt cropped.docx).
import re, sys, zipfile, shutil
from PIL import Image as PI, ImageDraw
im = PI.new("RGB", (900, 600), "white"); d = ImageDraw.Draw(im)
d.rectangle((0, 0, 199, 599), fill=(16, 124, 65)); d.rectangle((700, 0, 899, 599), fill=(16, 124, 65)); d.rectangle((0, 0, 899, 59), fill="black")
d.ellipse((250, 130, 650, 530), fill=(245, 150, 20)); im.save("sun.png")
SRC = '<a:srcRect xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" l="22222" t="10000" r="22222"/>' # openpyxl declares a: on each element
def patch(path, part):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            b = zi.read(i.filename)
            if re.fullmatch(part, i.filename): b = re.sub(rb'(<a:stretch)', SRC.encode() + rb'\1', b, count=1); assert SRC.encode() in b
            zo.writestr(i, b)
    shutil.move(tmp, path)
from docx import Document; from docx.shared import Inches
doc = Document(); doc.add_paragraph("A cropped picture: only the orange sun on white, no green, no black."); doc.add_picture("sun.png", width=Inches(2.5), height=Inches(2.7)); doc.add_paragraph("After the picture."); doc.save("cropped.docx"); patch("cropped.docx", r"word/document\.xml")
from openpyxl import Workbook; from openpyxl.drawing.image import Image
wb = Workbook(); ws = wb.active; ws["A1"] = "A cropped picture: only the sun"; img = Image("sun.png"); img.width, img.height = 250, 270; ws.add_image(img, "B3"); wb.save("cropped.xlsx"); patch("cropped.xlsx", r"xl/drawings/drawing1\.xml")
from pptx import Presentation; from pptx.util import Inches as I
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6]); pic = s.shapes.add_picture("sun.png", I(1), I(1), I(2.5), I(2.7))
pic.crop_left = pic.crop_right = 0.22222; pic.crop_top = 0.1; p.save("cropped.pptx")
