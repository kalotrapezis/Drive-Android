# Transitions and click animations test deck: slide 2 pushed in from below, slide 3 faded in with a box fading in on the
# first click and a star flying in from the left on the second — PowerPoint's own XML written into python-pptx's slides.
import sys, zipfile, shutil, re
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
out = sys.argv[1]
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = "Effects"; s.placeholders[1].text = "Tap to go on"
s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Pushed in from below"; s.placeholders[1].text = "A push transition"
s = p.slides.add_slide(p.slide_layouts[6])
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2)); box.text = "Fades in"
star = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(6), Inches(4), Inches(2.5), Inches(2.5))
ids = (box.shape_id, star.shape_id)
p.save(out)
def effect(n, spid, preset, sub, extra):
    return (f'<p:par><p:cTn id="{n}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{n+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:par><p:cTn id="{n+2}" presetID="{preset}" presetClass="entr" presetSubtype="{sub}" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{n+3}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'{extra}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
fade = lambda n, spid: f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{n}" dur="1000"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
fly = lambda n, spid: f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{n}" dur="1000" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>ppt_x</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:strVal val="0-#ppt_w/2"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#ppt_x"/></p:val></p:tav></p:tavLst></p:anim>'
timing = ('<p:timing><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
          + effect(3, ids[0], 10, 0, fade(7, ids[0])) + effect(10, ids[1], 2, 8, fly(14, ids[1])) +
          '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
add = {"ppt/slides/slide2.xml": '<p:transition spd="slow"><p:push dir="u"/></p:transition>', "ppt/slides/slide3.xml": '<p:transition spd="slow"><p:fade/></p:transition>' + timing}
tmp = out + ".tmp"; shutil.move(out, tmp)
with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        d = zi.read(it.filename)
        if it.filename in add: d = d.decode().replace("</p:clrMapOvr>", "</p:clrMapOvr>" + add[it.filename]).encode()
        zo.writestr(it, d)
import os; os.remove(tmp)
