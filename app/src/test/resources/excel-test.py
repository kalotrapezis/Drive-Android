# Excel-test.xlsx: one workbook, a sheet for each thing editing must keep right.
#  Rows   — a merged title, a table of numbers with formulas, a rule, a validation, a link, a comment, a chart, a picture
#  Refs   — formulas reading Rows (cells, ranges, whole columns)
#  Table  — an Excel Table (TableStyleMedium2) with a total formula
#  Formats — cells in many looks (fonts, fills, borders, alignments, number formats) for the formatting work
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image as PI
out = sys.argv[1]
wb = Workbook()
ws = wb.active; ws.title = "Rows"
ws["A1"] = "Grades 2026"; ws.merge_cells("A1:D1"); ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
ws.append(["Name", "Score", "Double", "Pass"])
names = ["Αλεξίου", "Βασιλείου", "Γεωργίου", "Δήμου", "Ελένη", "Ζαχαρίου", "Θεοδώρου", "Ιωάννου", "Κώστα", "Λάμπρου"]
for i, n in enumerate(names):
    r = i + 3
    ws.append([n, 10 + i * 1.5, f"=B{r}*2", f'=IF(B{r}>=15,"yes","no")'])
ws["A14"] = "Total"; ws["B14"] = "=SUM(B3:B12)"; ws["C14"] = "=SUM(C3:C12)"; ws["D14"] = '=COUNTIF(D3:D12,"yes")'
ws["A15"] = "Average"; ws["B15"] = "=AVERAGE(B$3:B$12)"
ws["A16"] = "Link"; ws["B16"].hyperlink = "#Refs!A1"; ws["B16"] = "to Refs"
ws["B5"].comment = Comment("Check this score", "Teo")
red = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
ws.conditional_formatting.add("B3:B12", CellIsRule(operator="greaterThan", formula=["20"], fill=red))
dv = DataValidation(type="list", formula1='"yes,no"'); ws.add_data_validation(dv); dv.add("D3:D12")
ch = BarChart(); ch.title = "Scores"; ch.add_data(Reference(ws, min_col=2, min_row=2, max_row=12), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=12)); ch.width = 10; ch.height = 6
ws.add_chart(ch, "F3")
im = PI.new("RGB", (60, 40), (16, 124, 65)); p = out.replace(".xlsx", ".png"); im.save(p); ws.add_image(Image(p), "F18")
ws.freeze_panes = "A3"
wb.defined_names["scores"] = __import__("openpyxl").workbook.defined_name.DefinedName("scores", attr_text="Rows!$B$3:$B$12")
ws.print_area = "A1:D16"

refs = wb.create_sheet("Refs")
refs["A1"] = "Sum of scores"; refs["B1"] = "=SUM(Rows!B3:B12)"
refs["A2"] = "Ελένη's score"; refs["B2"] = "=Rows!B7"
refs["A3"] = "Names"; refs["B3"] = "=COUNTA(Rows!A:A)"
refs["A4"] = "Lookup"; refs["B4"] = '=VLOOKUP("Δήμου",Rows!A3:B12,2,FALSE)'
refs["A5"] = "Named"; refs["B5"] = "=SUM(scores)"
refs["A6"] = "Own"; refs["B6"] = "=B1+B2"

t = wb.create_sheet("Table")
t.append(["Item", "Qty", "Price", "Cost"])
for i, (item, q, pr) in enumerate([("Pens", 10, 1.2), ("Paper", 5, 4.5), ("Ink", 2, 19.9), ("Tape", 7, 2.25), ("Clips", 30, 0.1)]):
    r = i + 2; t.append([item, q, pr, f"=B{r}*C{r}"])
t["C8"] = "Total"; t["D8"] = "=SUM(D2:D6)"
tab = Table(displayName="Supplies", ref="A1:D6"); tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); t.add_table(tab)

f = wb.create_sheet("Formats")
thin = Side(style="thin", color="FF000000"); thick = Side(style="thick", color="FF107C41")
f["A1"] = "Bold"; f["A1"].font = Font(bold=True)
f["A2"] = "Italic red"; f["A2"].font = Font(italic=True, color="FFFF0000")
f["A3"] = "Yellow fill"; f["A3"].fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
f["A4"] = "Boxed"; f["A4"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
f["A5"] = "Thick bottom"; f["A5"].border = Border(bottom=thick)
f["A6"] = "Centred"; f["A6"].alignment = Alignment(horizontal="center")
f["A7"] = "Wrapped text that runs on and on"; f["A7"].alignment = Alignment(wrap_text=True)
f["B1"] = 1234.5; f["B1"].number_format = "#,##0.00"
f["B2"] = 0.256; f["B2"].number_format = "0.0%"
f["B3"] = 46310; f["B3"].number_format = "dd/mm/yyyy"
f["B4"] = 1234.5; f["B4"].number_format = '#,##0.00 "€"'
f["B5"] = 12; f["B5"].font = Font(size=20, name="Arial")
f.column_dimensions["A"].width = 18; f.column_dimensions["B"].width = 14
wb.save(out)
