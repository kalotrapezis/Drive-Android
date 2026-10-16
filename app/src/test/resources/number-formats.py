import sys, datetime
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = "Formats"
codes = ["0", "0.00", "#,##0", "#,##0.00", "#,##0.00 \"€\"", "0%", "0.00%", "0.00E+00", "# ?/?", "# ??/??",
         "#,##0;[Red]-#,##0", "#,##0.00;-#,##0.00", "0.0;(0.0)", "dd/mm/yyyy", "d/m/yy", "dd mmm yyyy", "dddd d mmmm yyyy",
         "mmmm yyyy", "hh:mm", "h:mm:ss", "hh:mm AM/PM", "dd/mm/yyyy hh:mm", "000", "@", "\"Total \"0", "0.0#"]
values = [1234.567, -1234.567, 0.256, 0, 46310.5, 3.75, 12]
for r, code in enumerate(codes, start=1):
    ws.cell(r, 1, code)
    for c, v in enumerate(values, start=2):
        x = ws.cell(r, c, v); x.number_format = code
wb.save(sys.argv[1])
