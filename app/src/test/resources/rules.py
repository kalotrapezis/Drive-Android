# rules.xlsx: conditional formatting of each kind on a column of numbers A2:A11 (1..10) and words in B2:B11.
import sys
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, Rule
from openpyxl.styles import PatternFill, Font
from openpyxl.styles.differential import DifferentialStyle
out = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = "Rules"
ws.append(["n", "word", "scale", "bar", "stripe"])
words = ["apple", "Banana", "cherry", "apple", "date", "fig", "grape", "kiwi", "lemon", "Apple pie"]
for i in range(1, 11): ws.append([i, words[i - 1], i * 10, i, "x"])
red = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
green = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
ws.conditional_formatting.add("A2:A11", CellIsRule(operator="greaterThan", formula=["7"], fill=red, font=Font(color="FF9C0006")))
ws.conditional_formatting.add("A2:A11", CellIsRule(operator="between", formula=["2", "3"], fill=green))
r = Rule(type="containsText", operator="containsText", text="apple", dxf=DifferentialStyle(fill=green)); r.formula = ['NOT(ISERROR(SEARCH("apple",B2)))']
ws.conditional_formatting.add("B2:B11", r)
ws.conditional_formatting.add("C2:C11", ColorScaleRule(start_type="min", start_color="FFF8696B", end_type="max", end_color="FF63BE7B"))
ws.conditional_formatting.add("D2:D11", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
ws.conditional_formatting.add("E2:E11", FormulaRule(formula=["MOD(ROW(),2)=0"], fill=green))
wb.save(out)
