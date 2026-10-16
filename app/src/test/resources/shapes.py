# Shapes test workbook: Excel's shapes in their outlines — one in the theme's look (fillRef/lnRef, as Excel's default),
# others with their own colours, one with an adjustment (a thin arrow) and text. openpyxl makes the sheet; the drawing
# part is Excel's own XML, put in the file here.
import sys, zipfile, shutil
from openpyxl import Workbook
out = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = "Shapes"
ws["A1"] = "Shapes"; ws["A2"] = "Under the shapes"
wb.save(out)
def sp(i, prst, c1, r1, c2, r2, look, text="", av=""):
    tx = f'<xdr:txBody><a:bodyPr anchor="ctr"/><a:lstStyle/><a:p><a:pPr algn="ctr"/><a:r><a:t>{text}</a:t></a:r></a:p></xdr:txBody>' if text else ""
    return (f'<xdr:twoCellAnchor><xdr:from><xdr:col>{c1}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r1}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
            f'<xdr:to><xdr:col>{c2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>{r2}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{i + 1}" name="Shape {i}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst>{av}</a:avLst></a:prstGeom>{look}</xdr:spPr>'
            + ('' if '<a:solidFill' in look else '<xdr:style><a:lnRef idx="2"><a:schemeClr val="accent1"><a:shade val="50000"/></a:schemeClr></a:lnRef>'
               '<a:fillRef idx="1"><a:schemeClr val="accent1"/></a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent1"/></a:effectRef>'
               '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></xdr:style>')
            + f'{tx}</xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')
own = lambda fill, line: f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln w="19050"><a:solidFill><a:srgbClr val="{line}"/></a:solidFill></a:ln>'
shapes = [
    sp(1, "roundRect", 1, 3, 3, 7, "", "Theme look"),
    sp(2, "rightArrow", 4, 3, 6, 6, own("70AD47", "385723")),
    sp(3, "star5", 1, 9, 3, 14, own("FFC000", "BF9000")),
    sp(4, "heart", 4, 9, 6, 14, own("FF0000", "C00000")),
    sp(5, "ellipse", 7, 3, 9, 7, own("5B9BD5", "1F4E79"), "Oval"),
    sp(6, "rightArrow", 7, 9, 10, 12, own("C00000", "7F0000"), "", '<a:gd name="adj1" fmla="val 20000"/><a:gd name="adj2" fmla="val 70000"/>'),
]
drawing = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + "".join(shapes) + '</xdr:wsDr>')
tmp = out + ".tmp"; shutil.move(out, tmp)
with zipfile.ZipFile(tmp) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
    for item in zi.infolist():
        data = zi.read(item.filename)
        if item.filename == "[Content_Types].xml":
            data = data.replace(b"</Types>", b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
        if item.filename == "xl/worksheets/sheet1.xml":
            data = data.replace(b"</worksheet>", b'<drawing r:id="rId1" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/></worksheet>')
        zo.writestr(item, data)
    zo.writestr("xl/worksheets/_rels/sheet1.xml.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>')
    zo.writestr("xl/drawings/drawing1.xml", drawing)
import os; os.remove(tmp)
