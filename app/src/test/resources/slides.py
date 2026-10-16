# Slides test deck: a title slide, bullets in two levels (the layout's placeholders), shapes and a picture with a
# coloured text box, a chart, speaker notes — python-pptx's default template (its layouts, master and theme).
import sys, io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = "Tetra Slides"
s.placeholders[1].text = "A test deck, 16 Oct 2026"
s.notes_slide.notes_text_frame.text = "Say hello first."
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = "Plan"
tf = s.placeholders[1].text_frame
tf.text = "Read the slides"
for t, lvl in [("Text boxes and placeholders", 1), ("Pictures and shapes", 1), ("Present them", 0), ("Swipe to the next", 1)]:
    para = tf.add_paragraph(); para.text = t; para.level = lvl
s = p.slides.add_slide(p.slide_layouts[6])
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(0.5), Inches(4), Inches(1.5))
box.text = "A rounded box"
arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5), Inches(0.8), Inches(3), Inches(1))
arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor(0x70, 0xAD, 0x47)
tb = s.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(8), Inches(1))
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Bold red words, 28 pt"; r.font.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0xC0, 0, 0)
im = Image.new("RGB", (400, 240), (30, 120, 200))
for x in range(400):
    for y in range(100, 140): im.putpixel((x, y), (250, 200, 40))
buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
s.shapes.add_picture(buf, Inches(0.5), Inches(4), Inches(4), Inches(2.4))
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = "Scores"
cd = CategoryChartData(); cd.categories = ["Jan", "Feb", "Mar"]; cd.add_series("Sales", (120, 90, 150))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(5), cd)
p.save(sys.argv[1])
